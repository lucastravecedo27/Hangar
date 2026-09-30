"""Pruebas del módulo de Informes: catálogo, permisos por rol, las tres salidas de cada informe
(pantalla, Excel y PDF), entregables RAC con registros incompletos y que todo salga siempre al día.

    DATABASE_URL=sqlite:///copia.db python -m pytest pruebas/test_informes.py -q
"""
from datetime import date, timedelta
from io import BytesIO

import openpyxl
import pytest

from test_app import cliente, entrar  # noqa: F401  (fixture y ayuda compartidas)
import test_app
from core import db as core_db
from core import informes as core_inf
from core import informes_salida as core_sal

HOY = date.today()
MATRICULA = "HK-4321"


def _dias(n):
    return (HOY - timedelta(days=n)).isoformat()


def _con():
    return core_db.conectar(test_app.EMPRESA_A)


def _admin(c):
    return {"X-CSRF-Token": entrar(c, "prueba_admin", "clave-segura-1")}


@pytest.fixture(scope="module")
def avion(cliente):
    """Un avión con dos vuelos (uno completo y otro sin horas ni lugares), un trabajo de
    mantenimiento sin licencia del certificador y un piloto con licencia en Contactos."""
    h = _admin(cliente)
    con = _con()
    con.execute("DELETE FROM equipos WHERE matricula=?", (MATRICULA,))
    con.execute("DELETE FROM personal WHERE nombre='PILOTO INFORMES'")
    con.commit()
    r = cliente.post("/api/equipos", json={"nombre": f"Avión {MATRICULA}", "modelo": "C188A", "matricula": MATRICULA,
                                           "fecha_alta": _dias(90),
                                           "contadores": {"hobbs": 1000, "tach": 800, "aterrizajes": 2000,
                                                          "arranques": 900}}, headers=h)
    assert r.status_code == 201, r.data
    eid = r.get_json()["id"]
    pid = con.execute("INSERT INTO personal(nombre, rol, licencia, documento) VALUES('PILOTO INFORMES','piloto',"
                      "'PCA-1234','1010')").lastrowid
    con.commit()
    core_db.registrar_operacion(con, eid, 1.5, fecha=_dias(5), piloto="PILOTO INFORMES", personal_id=pid,
                                hora_salida="07:10", hora_llegada="08:45", origen="Pista La Ceiba",
                                destino="Pista La Ceiba", aterrizajes=6, hectareas=80, litros=800,
                                cliente="CLIENTE INFORMES", lote="Finca Uno", producto="Fungicida X")
    core_db.registrar_operacion(con, eid, 2.0, fecha=_dias(3), piloto="PILOTO INFORMES", personal_id=pid,
                                aterrizajes=8, hectareas=90, litros=900, cliente="CLIENTE INFORMES")
    core_db.registrar_mantenimiento(con, eid, None, "preventivo", "Inspección de 100 horas", 0, "Técnico Uno",
                                    fecha=_dias(2), firmado_por="Certificador Uno")
    con.commit()
    return eid


# ---------------------------------------------------------------- catálogo y permisos

def test_catalogo_admin_con_entregables_rac_primero(cliente, avion):
    _admin(cliente)
    r = cliente.get("/api/informes").get_json()
    claves = [i["clave"] for i in r["informes"]]
    for clave in ("rac_libro_abordo", "rac_libro_uas", "rac_aplicaciones", "rac_horas_pilotos", "rac_fiaa",
                  "rac_registro_mant", "resumen_gerencial", "operaciones_cliente", "estado_mantenimiento",
                  "ordenes_costos", "directivas_documentos", "rentabilidad", "combustible"):
        assert clave in claves, clave
    assert r["informes"][0]["categoria"] == "rac"
    assert r["categorias"][0]["clave"] == "rac"
    assert all(i["para"] for i in r["informes"])                    # cada informe dice para qué sirve
    assert all(i["exige"] for i in r["informes"] if i["categoria"] == "rac")
    assert "formato" in r["nota_rac"].lower()


def test_rentabilidad_solo_admin(cliente):
    entrar(cliente, "prueba_piloto", "clave-segura-3")
    claves = [i["clave"] for i in cliente.get("/api/informes").get_json()["informes"]]
    assert "rentabilidad" not in claves and "resumen_gerencial" not in claves and "ordenes_costos" not in claves
    assert "rac_libro_abordo" in claves
    assert cliente.get("/api/informes/rentabilidad").status_code == 403
    assert cliente.get("/informes/rentabilidad/excel").status_code == 403
    assert cliente.get("/informes/rentabilidad/pdf").status_code == 403
    entrar(cliente, "prueba_tecnico", "clave-segura-2")
    assert cliente.get("/api/informes/rentabilidad").status_code == 403
    assert cliente.get("/api/informes/ordenes_costos").status_code == 200


def test_sin_sesion_no_descarga(cliente):
    cliente.get("/salir")
    assert cliente.get("/informes/estado_mantenimiento/excel").status_code == 302
    assert cliente.get("/api/informes/estado_mantenimiento").status_code == 401


def test_informe_inexistente(cliente):
    _admin(cliente)
    assert cliente.get("/api/informes/no_existe").status_code == 404
    assert cliente.get("/informes/no_existe/pdf").status_code == 404


def test_pagina_y_enlace_directo(cliente):
    _admin(cliente)
    for url in ("/informes", "/informes?informe=rac_libro_abordo"):
        r = cliente.get(url)
        assert r.status_code == 200 and b"Informes" in r.data


# ---------------------------------------------------------------- las tres salidas de cada informe

def _claves_admin(cliente):
    _admin(cliente)
    return [i["clave"] for i in cliente.get("/api/informes").get_json()["informes"]]


def test_todos_los_informes_en_pantalla_excel_y_pdf(cliente, avion):
    for clave in _claves_admin(cliente):
        defn = core_inf.REGISTRO[clave]
        q = "periodo=todo" if "periodo" in defn["filtros"] else ""
        r = cliente.get(f"/api/informes/{clave}?{q}")
        assert r.status_code == 200, (clave, r.data[:300])
        d = r.get_json()
        assert d["titulo"] and d["secciones"] and d["generado"], clave
        for k in d["kpis"]:
            assert k["nota"], (clave, k["etiqueta"])                # todo KPI explica cómo se calcula
        # Excel: una hoja de resumen y una por sección, con la cabecera de marca y los filtros.
        x = cliente.get(f"/informes/{clave}/excel?{q}")
        assert x.status_code == 200, (clave, x.data[:300])
        assert "attachment" in x.headers["Content-Disposition"]
        wb = openpyxl.load_workbook(BytesIO(x.data))
        assert wb.sheetnames[0] == "Resumen" and len(wb.sheetnames) == 1 + len(d["secciones"]), clave
        ws = wb["Resumen"]
        assert ws["A1"].value.startswith("HANGAR") and ws["A2"].value == d["titulo"]
        assert "Generado el" in ws["A4"].value
        hoja = wb[wb.sheetnames[1]]
        fila_tit = 7 if hoja["A5"].value else 6
        titulos = [hoja.cell(row=fila_tit, column=i + 1).value for i in range(len(d["secciones"][0]["columnas"]))]
        assert titulos == [c["titulo"] for c in d["secciones"][0]["columnas"]], clave
        # PDF generado en el servidor
        p = cliente.get(f"/informes/{clave}/pdf?{q}")
        assert p.status_code == 200, (clave, p.data[:300])
        assert p.mimetype == "application/pdf" and p.data[:5] == b"%PDF-" and b"%%EOF" in p.data[-1024:], clave
        # Hoja de impresión
        assert cliente.get(f"/informes/{clave}/imprimir?{q}").status_code == 200, clave


def test_pdf_y_excel_validos_con_la_libreria(cliente, avion):
    """El PDF se abre con un lector (páginas y texto) y el Excel conserva formatos de número y fecha."""
    _admin(cliente)
    d = cliente.get("/api/informes/rac_libro_abordo?periodo=todo&equipo_id=%d" % avion).get_json()
    pdf = core_sal.pdf(d).getvalue()
    assert pdf.startswith(b"%PDF-") and b"/Type /Page" in pdf
    wb = openpyxl.load_workbook(core_sal.excel(d))
    ws = wb[wb.sheetnames[1]]
    fila_tit = 7
    cols = [ws.cell(row=fila_tit, column=i).value for i in range(1, 15)]
    assert cols[0] == "Fecha" and "Firma piloto al mando" in cols
    fecha = ws.cell(row=fila_tit + 1, column=1)
    assert fecha.is_date and fecha.number_format == "dd/mm/yyyy"
    horas = ws.cell(row=fila_tit + 1, column=cols.index("Tiempo de vuelo (h)") + 1)
    assert horas.number_format.startswith("#,##0")


# ---------------------------------------------------------------- entregables RAC

def test_libro_de_a_bordo_marca_incompletos(cliente, avion):
    _admin(cliente)
    d = cliente.get(f"/api/informes/rac_libro_abordo?periodo=30d&equipo_id={avion}").get_json()
    s = d["secciones"][0]
    assert d["norma"] == "RAC 91.1410" and d["conservar"] == "3 años"
    assert ("Matrícula", MATRICULA) in [tuple(x) for x in s["encabezado"]]
    assert ("Nacionalidad", "Colombia (HK)") in [tuple(x) for x in s["encabezado"]]
    assert len(s["filas"]) == 2
    completa = next(f for f in s["filas"] if f["hora_salida"])
    assert completa["licencia"] == "PCA-1234" and completa["horas_bloque"] == 1.58 and "_incompleto" not in completa
    assert completa["proposito"] == "Aplicación aérea"
    incompleta = next(f for f in s["filas"] if not f["hora_salida"])
    assert set(incompleta["_incompleto"]) == {"origen", "hora_salida", "destino", "hora_llegada"}
    assert any("1 de 2 registros incompletos" in a["texto"] for a in d["avisos"])
    assert next(k for k in d["kpis"] if k["etiqueta"] == "Registros incompletos")["valor"] == 1
    assert d["nota_rac"] and "formato de archivo" in d["nota_rac"]


def test_registro_de_mantenimiento_exige_licencia_del_certificador(cliente, avion):
    _admin(cliente)
    d = cliente.get(f"/api/informes/rac_registro_mant?periodo=30d&equipo_id={avion}").get_json()
    f = d["secciones"][0]["filas"][0]
    assert f["certifico"] == "Certificador Uno" and f["realizado"] == "Técnico Uno"
    assert f["_incompleto"] == ["licencia"]
    assert f["ciclos"] == 2014                        # 2000 iniciales + 6 + 8 aterrizajes


def test_registro_de_aplicaciones_137_71(cliente, avion):
    _admin(cliente)
    d = cliente.get("/api/informes/rac_aplicaciones?periodo=30d&cliente=CLIENTE%20INFORMES").get_json()
    filas = d["secciones"][0]["filas"]
    assert len(filas) == 2 and {f["cliente"] for f in filas} == {"CLIENTE INFORMES"}
    con_producto = next(f for f in filas if f["producto"])
    assert con_producto["cantidad"] == 800 and con_producto["dosis"] == 10.0
    assert "direccion" in con_producto["_incompleto"]           # el cliente no tiene dirección en Contactos
    sin = next(f for f in filas if not f["producto"])
    assert "producto" in sin["_incompleto"]


def test_certificacion_de_horas_por_piloto(cliente, avion):
    _admin(cliente)
    d = cliente.get("/api/informes/rac_horas_pilotos?periodo=30d&piloto=PILOTO%20INFORMES").get_json()
    res = d["secciones"][0]["filas"]
    assert len(res) == 1 and res[0]["horas"] == 3.5 and res[0]["licencia"] == "PCA-1234"
    assert res[0]["estado"] == "falta"                               # sin médico ni colinesterasa registrados
    assert d["secciones"][1]["firmas"]                               # firma del explotador por piloto


def test_fiaa_con_estadistica_directivas_y_componentes(cliente, avion):
    _admin(cliente)
    d = cliente.get(f"/api/informes/rac_fiaa?periodo=365d&equipo_id={avion}").get_json()
    titulos = [s["titulo"] for s in d["secciones"]]
    assert any(t.startswith("Estadística") for t in titulos)
    assert any(t.startswith("Directivas") for t in titulos)
    assert any(t.startswith("Componentes") for t in titulos)
    est = d["secciones"][0]
    assert est["totales"]["aterrizajes"] == 14 and est["totales"]["servicios"] == 2


# ---------------------------------------------------------------- siempre al día y filtros

def test_informe_refleja_lo_ultimo_registrado(cliente, avion):
    _admin(cliente)
    url = f"/api/informes/rac_libro_abordo?periodo=30d&equipo_id={avion}"
    antes = len(cliente.get(url).get_json()["secciones"][0]["filas"])
    con = _con()
    oid = core_db.registrar_operacion(con, avion, 0.5, fecha=_dias(1), piloto="PILOTO INFORMES")
    try:
        despues = cliente.get(url).get_json()["secciones"][0]["filas"]
        assert len(despues) == antes + 1
        x = cliente.get(f"/informes/rac_libro_abordo/excel?periodo=30d&equipo_id={avion}")
        ws = openpyxl.load_workbook(BytesIO(x.data)).worksheets[1]
        fechas = [ws.cell(row=r, column=1).value for r in range(8, 8 + antes + 1)]
        assert any(f and f.date().isoformat() == _dias(1) for f in fechas)
    finally:
        core_db.eliminar_operacion(con, oid)
        con.commit()


def test_filtro_por_tipo_de_activo(cliente, avion):
    _admin(cliente)
    d = cliente.get("/api/informes/productividad_activo?periodo=todo&tipo=avion").get_json()
    assert {f["tipo"] for f in d["secciones"][0]["filas"]} <= {"Avión"}
    d = cliente.get("/api/informes/productividad_activo?periodo=todo&tipo=dron").get_json()
    assert {f["tipo"] for f in d["secciones"][0]["filas"]} <= {"Dron"}
    assert "Drones" in d["filtros_txt"]


def test_parametros_invalidos_no_rompen(cliente):
    _admin(cliente)
    r = cliente.get("/api/informes/operaciones_cliente?desde=xx&hasta=2026-13-45&tipo=otro&equipo_id=abc&periodo=zz")
    assert r.status_code == 200
    d = r.get_json()
    assert d["filtros"]["tipo"] is None and d["filtros"]["equipo_id"] is None
    r = cliente.get("/api/informes/operaciones_cliente?cliente=%27%20OR%201=1--")
    assert r.status_code == 200 and not r.get_json()["secciones"][0]["filas"]


# ---------------------------------------------------------------- fijados por usuario

def test_fijar_informe_por_usuario(cliente):
    h = {"X-CSRF-Token": entrar(cliente, "prueba_piloto", "clave-segura-3")}
    r = cliente.put("/api/cuenta/informes-fijados", json={"clave": "rac_libro_uas", "fijado": True}, headers=h)
    assert r.status_code == 200 and r.get_json()["fijados"][0] == "rac_libro_uas"
    assert "rac_libro_uas" in cliente.get("/api/informes").get_json()["fijados"]
    # Un informe que su rol no puede ver no se fija.
    assert cliente.put("/api/cuenta/informes-fijados", json={"clave": "rentabilidad"}, headers=h).status_code == 400
    # Los fijados son de cada usuario.
    _admin(cliente)
    assert "rac_libro_uas" not in cliente.get("/api/informes").get_json()["fijados"]
    h = {"X-CSRF-Token": entrar(cliente, "prueba_piloto", "clave-segura-3")}
    cliente.put("/api/cuenta/informes-fijados", json={"clave": "rac_libro_uas", "fijado": False}, headers=h)
    assert "rac_libro_uas" not in cliente.get("/api/informes").get_json()["fijados"]


def test_fijar_exige_csrf(cliente):
    entrar(cliente, "prueba_piloto", "clave-segura-3")
    assert cliente.put("/api/cuenta/informes-fijados", json={"clave": "rac_libro_uas"}).status_code == 403


def test_limpieza(cliente, avion):
    con = _con()
    con.execute("DELETE FROM equipos WHERE matricula=?", (MATRICULA,))
    con.execute("DELETE FROM personal WHERE nombre='PILOTO INFORMES'")
    con.commit()
