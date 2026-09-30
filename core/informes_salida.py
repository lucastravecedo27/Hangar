"""Salidas de un informe (core/informes.py): Excel con openpyxl y PDF con ReportLab.

Las dos parten de la misma estructura que se ve en pantalla, así que el archivo descargado dice
exactamente lo mismo que la vista previa. Mismo aspecto de marca que las demás descargas de la app
(core/excel.py): banda grafito con la marca, título, filtros aplicados y fecha de generación.

PDF: se genera en el servidor con ReportLab (puro Python, se instala con pip en el Mac y en la
imagen de Docker, sin librerías del sistema). Si en alguna instalación no estuviera, la ruta de
descarga cae a la hoja de impresión HTML (templates/informe_imprimir.html) y el navegador la guarda
como PDF.
"""
from datetime import date, datetime
from io import BytesIO

from core import excel as core_excel
from core.informes import ESTADOS

try:                                   # ReportLab: PDF del lado del servidor
    import reportlab  # noqa: F401
    HAY_PDF = True
except ImportError:                    # pragma: no cover
    HAY_PDF = False

MARCA, LEMA = core_excel.MARCA, core_excel.LEMA
GRAFITO, GRAFITO_MED, SUAVE = core_excel.AZUL, core_excel.AZUL_MED, core_excel.AZUL_SUAVE
TINTA, GRIS, LINEA, BANDA = core_excel.TINTA, core_excel.GRIS, core_excel.LINEA, core_excel.BANDA
ACENTO = "FDB913"
# Semáforo (nunca el acento de marca): el mismo de las gráficas de la app. El ámbar se oscurece en
# el texto para que se lea sobre blanco.
SEMAFORO = {"ok": "2F8F4E", "proximo": "A86F00", "critico": "E0473E", "vencido": "B3261E", "gris": "737A86"}
FONDO_INCOMPLETO = "FDF0D5"
FONDO_AVISO = {"proximo": "FDF3DC", "critico": "FBE4E2", "vencido": "FBE4E2", "info": "EEF0F3", "ok": "E6F3EA"}


# ======================================================================
# Formato de valores (Colombia: punto de miles, coma decimal)
# ======================================================================

def _num(v, dec=0):
    s = f"{float(v):,.{dec}f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


def texto(v, c):
    """Valor de una celda como texto, igual que en pantalla."""
    if v is None or v == "":
        return ""
    t = c.get("tipo", "texto")
    try:
        if t == "fecha":
            return date.fromisoformat(str(v)[:10]).strftime("%d/%m/%Y")
        if t == "entero":
            return _num(v, 0)
        if t in ("numero", "horas"):
            return _num(v, c.get("dec", 1))
        if t == "dinero":
            return "$ " + _num(v, 0)
        if t == "pct":
            return _num(v, c.get("dec", 1)) + " %"
    except (TypeError, ValueError):
        return str(v)
    if t == "estado":
        return ESTADOS.get(v, (str(v), "gris"))[0]
    if t == "firma":
        return ""
    return str(v)


def texto_kpi(k):
    if k.get("valor") is None:
        return "—"
    c = {"tipo": k.get("tipo", "numero"), "dec": k.get("dec", 0)}
    v = texto(k["valor"], c)
    return f"{v} {k['unidad']}".strip() if k.get("unidad") else v


def nombre_archivo(r, ext):
    base = f"Hangar - {r['titulo']} - {date.today().isoformat()}"
    limpio = "".join(ch for ch in base if ch not in '\\/:*?"<>|')
    return f"{limpio}.{ext}"


def _meta(r):
    partes = [r["filtros_txt"]]
    if r.get("empresa"):
        partes.append(r["empresa"])
    partes.append("Generado el " + datetime.strptime(r["generado"], "%Y-%m-%d %H:%M").strftime("%d/%m/%Y %H:%M"))
    if r.get("usuario"):
        partes.append(f"por {r['usuario']}")
    return "  ·  ".join(partes)


def _columnas_visibles(s, pdf=False):
    return [c for c in s["columnas"]]


# ======================================================================
# Excel
# ======================================================================

def _formato_excel(c):
    t, d = c.get("tipo"), c.get("dec", 0)
    if t == "entero":
        return "#,##0"
    if t in ("numero", "horas"):
        return "#,##0" + ("." + "0" * d if d else "")
    if t == "dinero":
        return '"$" #,##0'
    if t == "pct":
        return "0" + ("." + "0" * d if d else "") + '" %"'
    if t == "fecha":
        return "dd/mm/yyyy"
    return None


def _ancho_excel(c, filas):
    if c.get("ancho"):
        return c["ancho"] + 2
    t = c.get("tipo")
    base = {"fecha": 12, "entero": 10, "numero": 12, "horas": 11, "dinero": 15, "pct": 10, "estado": 18,
            "firma": 26}.get(t)
    if base:
        return max(base, min(len(c["titulo"]) + 3, 22))
    largo = max([len(str(f.get(c["clave"]) or "")) for f in filas[:300]] + [len(c["titulo"]) + 2])
    return max(10, min(largo + 2, 48))


def _banda(ws, fila, n, valor, relleno, fuente, alto, borde_inferior=False):
    from openpyxl.styles import Alignment, Border, PatternFill, Side
    ultima = ws.cell(row=fila, column=max(n, 1)).column_letter
    if n > 1:
        ws.merge_cells(f"A{fila}:{ultima}{fila}")
    c = ws[f"A{fila}"]
    c.value = valor
    c.fill = PatternFill("solid", fgColor=relleno)
    c.font = fuente
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    if borde_inferior:
        c.border = Border(bottom=Side(style="medium", color=GRAFITO_MED))
    ws.row_dimensions[fila].height = alto


def _cabecera_excel(ws, n, titulo, subtitulo, meta):
    from openpyxl.styles import Font
    _banda(ws, 1, n, f"{MARCA}   ·   {LEMA}", GRAFITO, Font(color="FFFFFF", bold=True, size=15), 30)
    _banda(ws, 2, n, titulo, SUAVE, Font(color=GRAFITO, bold=True, size=12.5), 22)
    _banda(ws, 3, n, subtitulo, SUAVE, Font(color=TINTA, size=10.5), 30 if len(subtitulo or "") > 110 else 18)
    _banda(ws, 4, n, meta, SUAVE, Font(color=GRIS, size=9.5, italic=True), 16, borde_inferior=True)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GRAFITO


def _imprimir_excel(ws, fila_titulos, apaisado=True):
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.page_setup.orientation = "landscape" if apaisado else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if fila_titulos:
        ws.print_title_rows = f"1:{fila_titulos}"
    ws.oddFooter.left.text = f"{MARCA} · {LEMA}"
    ws.oddFooter.right.text = "Página &P de &N"
    ws.oddFooter.left.size = ws.oddFooter.right.size = 8


def _nombre_hoja(titulo, usados):
    limpio = "".join(ch for ch in titulo if ch not in "[]:*?/\\").replace("·", "-").strip()[:31] or "Hoja"
    nombre, i = limpio, 2
    while nombre.lower() in usados:
        sufijo = f" ({i})"
        nombre = limpio[:31 - len(sufijo)] + sufijo
        i += 1
    usados.add(nombre.lower())
    return nombre


def _firmas_excel(ws, fila, firmas, n):
    from openpyxl.styles import Border, Font, Side
    if not firmas:
        return fila
    fila += 2
    linea = Border(top=Side(style="thin", color=TINTA))
    col = 1
    paso = max(2, n // max(len(firmas), 1))
    for texto_firma in firmas:
        for k in range(col, min(col + paso - 1, n) + 1):
            ws.cell(row=fila + 2, column=k).border = linea
        c = ws.cell(row=fila + 3, column=col, value=texto_firma)
        c.font = Font(size=9.5, color=GRIS)
        col += paso
        if col > n:
            col, fila = 1, fila + 5
    return fila + 4


def excel(r):
    """El informe en un libro de Excel: hoja «Resumen» y una hoja por sección."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    usados = {"resumen"}
    ws.title = "Resumen"
    n = 5
    for letra, ancho in zip("ABCDE", (34, 20, 12, 22, 70)):
        ws.column_dimensions[letra].width = ancho
    _cabecera_excel(ws, n, r["titulo"], r.get("para") or r.get("subtitulo") or "", _meta(r))
    borde = core_excel._borde()
    fila = 6

    def titulo_bloque(texto_t):
        nonlocal fila
        fila += 1
        c = ws.cell(row=fila, column=1, value=texto_t.upper())
        c.font = Font(bold=True, size=9.5, color=GRAFITO_MED)
        c.border = Border(bottom=Side(style="thin", color=LINEA))
        fila += 1

    def par(etiqueta, valor, fuente=None):
        nonlocal fila
        ws.cell(row=fila, column=1, value=etiqueta).font = Font(size=10, color=GRIS, bold=True)
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=n)
        c = ws.cell(row=fila, column=2, value=valor)
        c.font = fuente or Font(size=10.5, color=TINTA)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        largo = len(str(valor or ""))
        ws.row_dimensions[fila].height = 15 * max(1, -(-largo // 120))
        fila += 1

    if r.get("subtitulo") and r.get("para"):
        par("Descripción", r["subtitulo"])
    if r.get("norma"):
        par("Norma", r["norma"] + (f"  ·  Conservar: {r['conservar']}" if r.get("conservar") else ""))
    if r.get("exige"):
        par("Qué exige", r["exige"])
    if r.get("nota_rac"):
        par("Formato", r["nota_rac"], Font(size=10, italic=True, color=GRIS))
    for a in r.get("avisos", []):
        par("Aviso" if a.get("nivel") != "info" else "Nota", a["texto"],
            Font(size=10.5, bold=a.get("nivel") != "info", color=SEMAFORO.get(a.get("nivel"), TINTA)))
        ws.cell(row=fila - 1, column=2).fill = PatternFill("solid", fgColor=FONDO_AVISO.get(a.get("nivel"), "EEF0F3"))
    if r.get("encabezado"):
        titulo_bloque("Datos del registro")
        for et, v in r["encabezado"]:
            par(et, v)

    if r.get("kpis"):
        titulo_bloque("Indicadores")
        for i, t in enumerate(("Indicador", "Valor", "Unidad", "Frente al período anterior", "Cómo se calcula"), 1):
            c = ws.cell(row=fila, column=i, value=t)
            c.fill = PatternFill("solid", fgColor=GRAFITO)
            c.font = Font(color="FFFFFF", bold=True, size=10)
            c.border = core_excel._borde("FFFFFF")
        fila += 1
        for k in r["kpis"]:
            ws.cell(row=fila, column=1, value=k["etiqueta"]).font = Font(size=10.5, color=TINTA, bold=True)
            cv = ws.cell(row=fila, column=2, value=k["valor"])
            cv.number_format = _formato_excel({"tipo": k.get("tipo"), "dec": k.get("dec", 0)}) or "General"
            cv.font = Font(size=11, bold=True, color=SEMAFORO.get(ESTADOS.get(k.get("estado"), ("", ""))[1], TINTA)
                           if k.get("estado") and k.get("estado") != "ok" else TINTA)
            ws.cell(row=fila, column=3, value=k.get("unidad") or None)
            if k.get("variacion") is not None:
                ws.cell(row=fila, column=4, value=f"{k['variacion']:+.1f} %".replace(".", ","))
            c = ws.cell(row=fila, column=5, value=k.get("nota"))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = Font(size=9.5, color=GRIS)
            for i in range(1, 6):
                ws.cell(row=fila, column=i).border = borde
            fila += 1

    titulo_bloque("Contenido")
    for i, t in enumerate(("Hoja", "Registros", "", "", "Descripción"), 1):
        c = ws.cell(row=fila, column=i, value=t or None)
        c.fill = PatternFill("solid", fgColor=GRAFITO)
        c.font = Font(color="FFFFFF", bold=True, size=10)
    fila += 1
    hojas = []
    for s in r["secciones"]:
        nombre = _nombre_hoja(s["titulo"], usados)
        hojas.append(nombre)
        c = ws.cell(row=fila, column=1, value=s["titulo"])
        c.hyperlink = f"#'{nombre}'!A1"
        c.font = Font(color="3B6FB6", underline="single", size=10.5)
        ws.cell(row=fila, column=2, value=len(s["filas"])).number_format = "#,##0"
        d = ws.cell(row=fila, column=5, value=s.get("descripcion"))
        d.font = Font(size=9.5, color=GRIS)
        fila += 1
    _firmas_excel(ws, fila, r.get("firmas"), n)
    _imprimir_excel(ws, None, apaisado=False)

    for s, nombre in zip(r["secciones"], hojas):
        _hoja_seccion(wb.create_sheet(nombre), s, r)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _hoja_seccion(ws, s, r):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    cols = _columnas_visibles(s)
    n = max(len(cols), 3)
    sub = r["titulo"] + (f"  ·  {r['norma']}" if r.get("norma") else "")
    _cabecera_excel(ws, n, s["titulo"], sub, _meta(r))
    extra = []
    if s.get("encabezado"):
        extra.append("   ·   ".join(f"{e}: {v}" for e, v in s["encabezado"]))
    if s.get("descripcion"):
        extra.append(s["descripcion"])
    ftit = 6
    if extra:
        _banda(ws, 5, n, "\n".join(extra), "FFFFFF", Font(size=10, color=TINTA, bold=bool(s.get("encabezado"))),
               16 * len(extra) + (8 if len(extra[0]) > 150 else 0))
        ftit = 7
    else:
        ws.row_dimensions[5].height = 7

    for i, c in enumerate(cols, 1):
        celda = ws.cell(row=ftit, column=i, value=c["titulo"])
        celda.fill = PatternFill("solid", fgColor=GRAFITO)
        celda.font = Font(color="FFFFFF", bold=True, size=10)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = core_excel._borde("FFFFFF")
        ws.column_dimensions[celda.column_letter].width = _ancho_excel(c, s["filas"])
    ws.row_dimensions[ftit].height = 32

    borde = core_excel._borde()
    banda = PatternFill("solid", fgColor=BANDA)
    incompleto = PatternFill("solid", fgColor=FONDO_INCOMPLETO)
    fuente = Font(size=10, color=TINTA)
    fuentes_estado = {k: Font(size=10, bold=True, color=v) for k, v in SEMAFORO.items()}
    formatos = [_formato_excel(c) for c in cols]
    envolver = Alignment(wrap_text=True, vertical="top")
    arriba = Alignment(vertical="top")
    fila = ftit + 1
    for i_f, f in enumerate(s["filas"]):
        faltan = set(f.get("_incompleto") or [])
        for i, c in enumerate(cols, 1):
            v = f.get(c["clave"])
            celda = ws.cell(row=fila, column=i)
            if c["tipo"] == "fecha" and v:
                try:
                    v = date.fromisoformat(str(v)[:10])
                except ValueError:
                    pass
            elif c["tipo"] == "estado" and v:
                txt, clase = ESTADOS.get(v, (str(v), "gris"))
                v = txt
                celda.font = fuentes_estado.get(clase, fuente)
            elif c["tipo"] == "firma":
                v = None
            celda.value = v
            if c["tipo"] != "estado" or not f.get(c["clave"]):
                celda.font = fuente
            if formatos[i - 1]:
                celda.number_format = formatos[i - 1]
            celda.border = borde
            celda.alignment = envolver if c["tipo"] == "texto" and (c.get("ancho") or 0) >= 24 else arriba
            if c["clave"] in faltan:
                celda.fill = incompleto
            elif i_f % 2:
                celda.fill = banda
        if any(c["tipo"] == "firma" for c in cols):
            ws.row_dimensions[fila].height = 28
        fila += 1
    if not s["filas"]:
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n)
        c = ws.cell(row=fila, column=1, value=s.get("vacio"))
        c.font = Font(size=10, italic=True, color=GRIS)
        fila += 1
    elif s.get("totales"):
        for i, c in enumerate(cols, 1):
            celda = ws.cell(row=fila, column=i)
            if i == 1 and c["clave"] not in s["totales"]:
                celda.value = "Total"
            elif c["clave"] in s["totales"]:
                celda.value = s["totales"][c["clave"]]
                if formatos[i - 1]:
                    celda.number_format = formatos[i - 1]
            celda.font = Font(size=10, bold=True, color=GRAFITO)
            celda.fill = PatternFill("solid", fgColor=SUAVE)
            celda.border = Border(top=Side(style="medium", color=GRAFITO_MED), bottom=Side(style="thin", color=LINEA))
        fila += 1
    if s["filas"]:
        ultima = ws.cell(row=ftit, column=len(cols)).column_letter
        ws.auto_filter.ref = f"A{ftit}:{ultima}{ftit + len(s['filas'])}"
    ws.freeze_panes = f"A{ftit + 1}"
    _firmas_excel(ws, fila, s.get("firmas"), n)
    _imprimir_excel(ws, ftit, apaisado=len(cols) > 6)


# ======================================================================
# PDF (ReportLab)
# ======================================================================

# Helvetica (fuente base de PDF) cubre el alfabeto latino (Windows-1252). Lo que se sale de ahí se
# escribe con equivalentes para no imprimir cuadros vacíos.
_SUSTITUTOS = {"≥": ">=", "≤": "<=", "→": "->", "←": "<-", "Σ": "Suma ", "−": "-", "✓": "Sí", "✕": "x",
               "÷": "/", "≈": "~", "▲": "", "▼": "", "↕": "", " ": " ", " ": " "}


def _limpio(t):
    if t is None:
        return ""
    salida = []
    for ch in str(t):
        try:
            ch.encode("cp1252")
            salida.append(ch)
        except UnicodeEncodeError:
            salida.append(_SUSTITUTOS.get(ch, "?"))
    return "".join(salida)


def _esc(t):
    return _limpio(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _color(h):
    from reportlab.lib import colors
    return colors.HexColor("#" + h.lstrip("#"))


def _dibujar_marca(c, x, y, tam):
    """La marca de Hangar (static_shell/marca.svg) dibujada con trazos: tres barras inclinadas."""
    import math
    k = tam / 120.0
    c.saveState()
    c.translate(x, y + tam)
    c.scale(k, -k)
    c.saveState()
    c.transform(1, 0, math.tan(math.radians(-14)), 1, 0, 0)
    c.setFillColor(_color(GRAFITO))
    c.roundRect(34, 66, 16, 32, 4, stroke=0, fill=1)
    c.roundRect(60, 44, 16, 54, 4, stroke=0, fill=1)
    c.setFillColor(_color(ACENTO))
    p = c.beginPath()
    p.moveTo(86, 36)
    p.curveTo(86, 26, 94, 14, 94, 14)
    p.curveTo(94, 14, 102, 26, 102, 36)
    p.lineTo(102, 94)
    p.curveTo(102, 96.2, 100.2, 98, 98, 98)
    p.lineTo(90, 98)
    p.curveTo(87.8, 98, 86, 96.2, 86, 94)
    p.close()
    c.drawPath(p, stroke=0, fill=1)
    c.restoreState()
    c.setFillColor(_color("B9BBC0"))
    c.roundRect(20, 104, 86, 6, 3, stroke=0, fill=1)
    c.restoreState()


def _estilos():
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    base = dict(fontName="Helvetica", textColor=_color(TINTA))
    return {
        "titulo": ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=_color(GRAFITO)),
        "sub": ParagraphStyle("sub", **base, fontSize=9.5, leading=12.5),
        "meta": ParagraphStyle("meta", fontName="Helvetica-Oblique", fontSize=8, leading=10.5, textColor=_color(GRIS)),
        "para": ParagraphStyle("para", fontName="Helvetica-Oblique", fontSize=9, leading=12, textColor=_color(GRAFITO_MED)),
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=11.5, leading=14, textColor=_color(GRAFITO),
                             spaceBefore=8, spaceAfter=3),
        "desc": ParagraphStyle("desc", **base, fontSize=8, leading=10),
        "celda": ParagraphStyle("celda", **base, fontSize=7.4, leading=9),
        "celda_num": ParagraphStyle("celda_num", **base, fontSize=7.4, leading=9, alignment=TA_RIGHT),
        "cab": ParagraphStyle("cab", fontName="Helvetica-Bold", fontSize=7.2, leading=8.6, textColor=_color("FFFFFF")),
        "cab_num": ParagraphStyle("cab_num", fontName="Helvetica-Bold", fontSize=7.2, leading=8.6,
                                  textColor=_color("FFFFFF"), alignment=TA_RIGHT),
        "kpi_l": ParagraphStyle("kpi_l", fontName="Helvetica-Bold", fontSize=6.8, leading=8.4, textColor=_color(GRIS)),
        "kpi_v": ParagraphStyle("kpi_v", fontName="Helvetica-Bold", fontSize=14, leading=17, textColor=_color(GRAFITO)),
        "kpi_n": ParagraphStyle("kpi_n", fontName="Helvetica", fontSize=6.5, leading=8, textColor=_color(GRIS)),
        "kv_l": ParagraphStyle("kv_l", fontName="Helvetica-Bold", fontSize=7, leading=9, textColor=_color(GRIS)),
        "kv_v": ParagraphStyle("kv_v", fontName="Helvetica-Bold", fontSize=8.8, leading=11, textColor=_color(TINTA)),
        "aviso": ParagraphStyle("aviso", **base, fontSize=8.5, leading=11),
        "firma": ParagraphStyle("firma", fontName="Helvetica", fontSize=8, leading=10, textColor=_color(GRIS)),
    }


_NUMERICOS = ("entero", "numero", "horas", "dinero", "pct")


def _anchos_pdf(cols, disponible):
    pesos = []
    for c in cols:
        t = c["tipo"]
        if c.get("ancho"):
            w = c["ancho"]
        else:
            w = {"fecha": 10, "entero": 8, "numero": 9, "horas": 9, "dinero": 11, "pct": 8, "estado": 12,
                 "firma": 22}.get(t, 14)
        pesos.append(max(w, 6))
    total = sum(pesos)
    return [disponible * w / total for w in pesos]


def _tabla_pdf(s, est, ancho):
    from reportlab.platypus import Paragraph, Table, TableStyle
    cols = _columnas_visibles(s, pdf=True)
    fila_cab = [Paragraph(_esc(c["titulo"]), est["cab_num"] if c["tipo"] in _NUMERICOS else est["cab"]) for c in cols]
    datos = [fila_cab]
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), _color(GRAFITO)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.3, _color("E6E7EA")),
        ("TOPPADDING", (0, 0), (-1, -1), 2.6), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.6),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    firma = any(c["tipo"] == "firma" for c in cols)
    for i, f in enumerate(s["filas"], 1):
        faltan = set(f.get("_incompleto") or [])
        fila = []
        for j, c in enumerate(cols):
            v = f.get(c["clave"])
            t = texto(v, c)
            if c["tipo"] == "estado" and v:
                clase = ESTADOS.get(v, ("", "gris"))[1]
                t = f'<font color="#{SEMAFORO.get(clase, GRIS)}"><b>{_esc(t)}</b></font>'
                fila.append(Paragraph(t, est["celda"]))
            else:
                fila.append(Paragraph(_esc(t), est["celda_num"] if c["tipo"] in _NUMERICOS else est["celda"]))
            if c["clave"] in faltan:
                estilo.append(("BACKGROUND", (j, i), (j, i), _color(FONDO_INCOMPLETO)))
        datos.append(fila)
        if i % 2 == 0:
            estilo.append(("BACKGROUND", (0, i), (-1, i), _color("F6F7F8")))
        if firma:
            estilo.append(("BOTTOMPADDING", (0, i), (-1, i), 12))
    if s.get("totales") and s["filas"]:
        tot = []
        for j, c in enumerate(cols):
            if c["clave"] in s["totales"]:
                tot.append(Paragraph("<b>" + _esc(texto(s["totales"][c["clave"]], c)) + "</b>", est["celda_num"]))
            else:
                tot.append(Paragraph("<b>Total</b>" if j == 0 else "", est["celda"]))
        datos.append(tot)
        n = len(datos) - 1
        estilo += [("BACKGROUND", (0, n), (-1, n), _color(SUAVE)), ("LINEABOVE", (0, n), (-1, n), 0.9, _color(GRAFITO_MED))]
    t = Table(datos, colWidths=_anchos_pdf(cols, ancho), repeatRows=1, splitByRow=1)
    t.setStyle(TableStyle(estilo))
    return t


def _kv_pdf(pares, est, ancho, por_fila=4):
    from reportlab.platypus import Paragraph, Table, TableStyle
    celdas = [[Paragraph(_esc(e).upper(), est["kv_l"]), Paragraph(_esc(v if v not in (None, "") else "—"), est["kv_v"])]
              for e, v in pares]
    filas = []
    for i in range(0, len(celdas), por_fila):
        trozo = celdas[i:i + por_fila]
        trozo += [["", ""]] * (por_fila - len(trozo))
        filas.append([x for par_ in trozo for x in par_])
    anchos = []
    for _ in range(por_fila):
        anchos += [ancho / por_fila * 0.42, ancho / por_fila * 0.58]
    t = Table(filas, colWidths=anchos)
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, -1), _color("F6F7F8")),
                           ("BOX", (0, 0), (-1, -1), 0.4, _color("E6E7EA")),
                           ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    return t


def _kpis_pdf(kpis, est, ancho, por_fila):
    from reportlab.platypus import Paragraph, Table, TableStyle
    celdas = []
    for k in kpis:
        color = GRAFITO
        if k.get("estado") and k["estado"] != "ok":
            color = SEMAFORO.get(ESTADOS.get(k["estado"], ("", "gris"))[1], GRAFITO)
        valor = f'<font color="#{color}">{_esc(texto_kpi(k))}</font>'
        nota = k.get("nota") or ""
        if k.get("variacion") is not None:
            nota = (f"{k['variacion']:+.1f} % frente al período anterior".replace(".", ",")
                    + (" · " + nota if nota else ""))
        celdas.append([Paragraph(_esc(k["etiqueta"]).upper(), est["kpi_l"]), Paragraph(valor, est["kpi_v"]),
                       Paragraph(_esc(nota), est["kpi_n"])])
    filas = []
    for i in range(0, len(celdas), por_fila):
        trozo = celdas[i:i + por_fila]
        trozo += [""] * (por_fila - len(trozo))
        filas.append(trozo)
    t = Table(filas, colWidths=[ancho / por_fila] * por_fila)
    estilo = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 6),
              ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LEFTPADDING", (0, 0), (-1, -1), 7)]
    for r_i, fila in enumerate(filas):
        for c_i, celda in enumerate(fila):
            if celda:
                estilo.append(("BOX", (c_i, r_i), (c_i, r_i), 0.5, _color("E6E7EA")))
    t.setStyle(TableStyle(estilo))
    return t


_PALETA = ["FDB913", "1A1D24", "0D9488", "9757C9", "3B6FB6", "2F8F4E", "737A86", "E0473E"]
_NIVEL_GRAF = {"ok": "2F8F4E", "proximo": "E8A400", "critico": "E0473E", "vencido": "B3261E", "gris": "9AA0A8"}


def _grafica_pdf(g, ancho, alto=150):
    """Barras (verticales u horizontales) con ReportLab. Las donas se dibujan como barras."""
    from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
    from reportlab.graphics.shapes import Drawing, String
    etiquetas = [_limpio(e)[:18] for e in g["etiquetas"]]
    series = [s for s in g["series"] if s.get("datos")]
    if not etiquetas or not series or not any(any(v for v in s["datos"]) for s in series):
        return None
    d = Drawing(ancho, alto)
    d.add(String(0, alto - 10, _limpio(g["titulo"]), fontName="Helvetica-Bold", fontSize=9, fillColor=_color(GRAFITO)))
    horizontal = g["tipo"] in ("ranking",) or (g["tipo"] == "dona")
    if horizontal:
        etiquetas, series = etiquetas[:10], [dict(s, datos=s["datos"][:10]) for s in series]
        ch = HorizontalBarChart()
        ch.x, ch.y, ch.width, ch.height = 90, 8, ancho - 110, alto - 30
        ch.categoryAxis.labels.fontSize = 6.5
        ch.categoryAxis.reverseDirection = 1
        ch.valueAxis.labels.fontSize = 6.5
    else:
        ch = VerticalBarChart()
        ch.x, ch.y, ch.width, ch.height = 32, 22, ancho - 42, alto - 44
        ch.categoryAxis.labels.fontSize = 6.2
        ch.categoryAxis.labels.angle = 30 if len(etiquetas) > 8 else 0
        ch.categoryAxis.labels.boxAnchor = "ne" if len(etiquetas) > 8 else "n"
        ch.valueAxis.labels.fontSize = 6.5
    ch.data = [[float(v or 0) for v in s["datos"]] for s in series]
    ch.categoryAxis.categoryNames = etiquetas
    ch.categoryAxis.labels.fontName = ch.valueAxis.labels.fontName = "Helvetica"
    ch.valueAxis.valueMin = 0
    ch.valueAxis.gridStrokeColor = _color("EDEDF0")
    ch.valueAxis.visibleGrid = 1
    ch.categoryAxis.strokeColor = ch.valueAxis.strokeColor = _color("D3D6DC")
    ch.barSpacing, ch.groupSpacing = 1, 4
    for i, s in enumerate(series):
        color = _NIVEL_GRAF.get(s.get("nivel")) or _PALETA[i % len(_PALETA)]
        ch.bars[i].fillColor = _color(color)
        ch.bars[i].strokeColor = None
    if g["tipo"] == "dona" and g.get("colores"):
        for j, clase in enumerate(g["colores"][:len(etiquetas)]):
            ch.bars[(0, j)].fillColor = _color(_NIVEL_GRAF.get(clase, "9AA0A8"))
    d.add(ch)
    if len(series) > 1:
        x = ancho - 10
        for i, s in reversed(list(enumerate(series))):
            nombre = _limpio(s["nombre"])
            x -= 8 + len(nombre) * 3.6
            color = _NIVEL_GRAF.get(s.get("nivel")) or _PALETA[i % len(_PALETA)]
            from reportlab.graphics.shapes import Rect
            d.add(Rect(x, alto - 10, 6, 6, fillColor=_color(color), strokeColor=None))
            d.add(String(x + 8, alto - 9.5, nombre, fontName="Helvetica", fontSize=6.5, fillColor=_color(GRIS)))
    return d


def _firmas_pdf(firmas, est, ancho):
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
    if not firmas:
        return []
    por_fila = 2 if len(firmas) != 3 else 3
    celdas = [[Spacer(1, 26), Paragraph(_esc(f), est["firma"])] for f in firmas]
    filas = []
    for i in range(0, len(celdas), por_fila):
        trozo = celdas[i:i + por_fila]
        filas.append([Table([[c[0]], [c[1]]], colWidths=[ancho / por_fila - 24],
                            style=TableStyle([("LINEBELOW", (0, 0), (0, 0), 0.7, _color(TINTA)),
                                              ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
                      for c in trozo] + [""] * (por_fila - len(trozo)))
    return [Spacer(1, 14), Table(filas, colWidths=[ancho / por_fila] * por_fila,
                                 style=TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0),
                                                   ("TOPPADDING", (0, 0), (-1, -1), 8)]))]


def pdf(r):
    """El informe en PDF: cabecera de marca en cada página, pie con paginación «Página X de Y»."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import (BaseDocTemplate, CondPageBreak, Frame, KeepTogether, PageTemplate,
                                    Paragraph, Spacer, Table, TableStyle)

    columnas = max([len(s["columnas"]) for s in r["secciones"]] + [0])
    tam = landscape(A4) if columnas > 7 else A4
    margen = 13 * mm
    ancho = tam[0] - 2 * margen
    est = _estilos()
    generado = datetime.strptime(r["generado"], "%Y-%m-%d %H:%M").strftime("%d/%m/%Y %H:%M")
    pie_izq = _limpio(f"{MARCA.capitalize()} · {LEMA}" + (f" · {r['empresa']}" if r.get("empresa") else ""))
    titulo_corto = _limpio(r["titulo"])

    class Lienzo(rl_canvas.Canvas):
        """Guarda las páginas para escribir el total («de Y») al final."""

        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._paginas = []

        def showPage(self):
            self._paginas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._paginas)
            for estado in self._paginas:
                self.__dict__.update(estado)
                self._decorar(total)
                super().showPage()
            super().save()

        def _decorar(self, total):
            w, h = tam
            # Cabecera: marca, nombre y título del informe; filete amarillo de marca debajo.
            _dibujar_marca(self, margen, h - margen - 16, 16)
            self.setFont("Helvetica-Bold", 11)
            self.setFillColor(_color(GRAFITO))
            self.drawString(margen + 21, h - margen - 8.5, MARCA)
            self.setFont("Helvetica", 7)
            self.setFillColor(_color(GRIS))
            self.drawString(margen + 21, h - margen - 16, _limpio(LEMA))
            self.setFont("Helvetica-Bold", 8.5)
            self.setFillColor(_color(GRAFITO))
            self.drawRightString(w - margen, h - margen - 8.5, titulo_corto[:110])
            self.setFont("Helvetica", 7)
            self.setFillColor(_color(GRIS))
            self.drawRightString(w - margen, h - margen - 16, _limpio(r.get("norma") or r["categoria_txt"]))
            self.setStrokeColor(_color(ACENTO))
            self.setLineWidth(1.4)
            self.line(margen, h - margen - 21, w - margen, h - margen - 21)
            # Pie
            self.setStrokeColor(_color("E6E7EA"))
            self.setLineWidth(0.5)
            self.line(margen, margen - 2, w - margen, margen - 2)
            self.setFont("Helvetica", 7)
            self.setFillColor(_color(GRIS))
            self.drawString(margen, margen - 10, pie_izq)
            self.drawCentredString(w / 2, margen - 10, f"Generado el {generado}")
            self.drawRightString(w - margen, margen - 10, f"Página {self._pageNumber} de {total}")

    buffer = BytesIO()
    doc = BaseDocTemplate(buffer, pagesize=tam, leftMargin=margen, rightMargin=margen, topMargin=margen + 26,
                          bottomMargin=margen + 4, title=_limpio(r["titulo"]), author=_limpio(r.get("empresa") or MARCA),
                          subject=_limpio(r.get("norma") or r["categoria_txt"]), creator="Hangar")
    marco = Frame(margen, margen + 4, ancho, tam[1] - 2 * margen - 30, id="cuerpo", leftPadding=0, rightPadding=0,
                  topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id="pagina", frames=[marco])])

    h = [Paragraph(_esc(r["titulo"]), est["titulo"]), Spacer(1, 3)]
    if r.get("subtitulo"):
        h.append(Paragraph(_esc(r["subtitulo"]), est["sub"]))
    h.append(Paragraph(_esc(_meta(r)), est["meta"]))
    if r.get("para"):
        h += [Spacer(1, 3), Paragraph("<b>Para qué sirve:</b> " + _esc(r["para"]), est["para"])]
    h.append(Spacer(1, 8))

    cajas = []
    if r.get("norma") or r.get("exige"):
        texto_norma = ""
        if r.get("norma"):
            texto_norma += f"<b>{_esc(r['norma'])}</b>" + (f" · Conservar: {_esc(r['conservar'])}" if r.get("conservar") else "")
        if r.get("exige"):
            texto_norma += ("<br/>" if texto_norma else "") + "<b>Qué exige:</b> " + _esc(r["exige"])
        if r.get("nota_rac"):
            texto_norma += "<br/><i>" + _esc(r["nota_rac"]) + "</i>"
        cajas.append(("EEF0F3", GRAFITO_MED, texto_norma))
    for a in r.get("avisos", []):
        nivel = a.get("nivel", "info")
        cajas.append((FONDO_AVISO.get(nivel, "EEF0F3"), SEMAFORO.get(nivel if nivel in SEMAFORO else "gris", GRIS),
                      ("<b>Atención:</b> " if nivel != "info" else "") + _esc(a["texto"])))
    for fondo, borde, t in cajas:
        caja = Table([[Paragraph(t, est["aviso"])]], colWidths=[ancho])
        caja.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), _color(fondo)),
                                  ("LINEBEFORE", (0, 0), (0, -1), 2.2, _color(borde)),
                                  ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
        h += [caja, Spacer(1, 5)]

    if r.get("encabezado"):
        h += [_kv_pdf(r["encabezado"], est, ancho, 4 if tam[0] > tam[1] else 3), Spacer(1, 8)]
    if r.get("kpis"):
        h += [_kpis_pdf(r["kpis"], est, ancho, 5 if tam[0] > tam[1] else 4), Spacer(1, 8)]

    graficas = [g for g in (_grafica_pdf(g, ancho / 2 - 8) for g in r.get("graficas", [])[:2]) if g]
    if graficas:
        fila = graficas + [""] * (2 - len(graficas))
        h += [Table([fila], colWidths=[ancho / 2] * 2), Spacer(1, 6)]

    for s in r["secciones"]:
        bloque = [CondPageBreak(60 * mm), Paragraph(_esc(s["titulo"]), est["h2"])]
        if s.get("descripcion"):
            bloque.append(Paragraph(_esc(s["descripcion"]), est["desc"]))
        if s.get("encabezado"):
            bloque += [Spacer(1, 3), _kv_pdf(s["encabezado"], est, ancho, 4 if tam[0] > tam[1] else 3)]
        bloque.append(Spacer(1, 4))
        h += bloque
        if s["filas"]:
            h.append(_tabla_pdf(s, est, ancho))
        else:
            h.append(Paragraph("<i>" + _esc(s.get("vacio")) + "</i>", est["desc"]))
        if s.get("firmas"):
            h.append(KeepTogether(_firmas_pdf(s["firmas"], est, ancho)))
        h.append(Spacer(1, 6))
    if r.get("firmas"):
        h.append(KeepTogether([CondPageBreak(40 * mm)] + _firmas_pdf(r["firmas"], est, ancho)))

    doc.build(h, canvasmaker=Lienzo)
    buffer.seek(0)
    return buffer


# ======================================================================
# Hoja de impresión HTML (respaldo sin ReportLab)
# ======================================================================

def para_imprimir(r):
    """La estructura con cada celda ya formateada como texto, para la plantilla de impresión."""
    salida = dict(r)
    salida["kpis_txt"] = [dict(k, texto=texto_kpi(k), clase=ESTADOS.get(k.get("estado"), ("", ""))[1]) for k in r["kpis"]]
    secciones = []
    for s in r["secciones"]:
        cols = _columnas_visibles(s)
        filas = []
        for f in s["filas"]:
            faltan = set(f.get("_incompleto") or [])
            filas.append([{"t": texto(f.get(c["clave"]), c), "num": c["tipo"] in _NUMERICOS,
                           "clase": ESTADOS.get(f.get(c["clave"]), ("", ""))[1] if c["tipo"] == "estado" else "",
                           "falta": c["clave"] in faltan} for c in cols])
        tot = None
        if s.get("totales") and s["filas"]:
            tot = [{"t": texto(s["totales"].get(c["clave"]), c) if c["clave"] in s["totales"] else ("Total" if i == 0 else ""),
                    "num": c["tipo"] in _NUMERICOS} for i, c in enumerate(cols)]
        secciones.append(dict(s, cols=cols, filas_txt=filas, total_txt=tot))
    salida["secciones_txt"] = secciones
    salida["apaisado"] = max([len(s["columnas"]) for s in r["secciones"]] + [0]) > 7
    return salida
